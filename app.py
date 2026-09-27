import streamlit as st
import PyPDF2
import io
import zipfile
import base64
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

# 1. Page Configuration
st.set_page_config(
    page_title="Anuan Ratu Kaylla dabel L",
    page_icon="⚡",
    layout="wide"
)

# 2. Custom CSS for Light Modern UI
st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@400;500;600;700&display=swap');
    
    html, body, [class*="css"] {
        font-family: 'Plus Jakarta Sans', sans-serif !important;
    }
    
    .stApp {
        background-color: #F8FAFC !important;
    }

    /* Hero Banner */
    .hero-banner {
        background: linear-gradient(135deg, #EFF6FF 0%, #DBEAFE 100%);
        border: 1px solid #BFDBFE;
        border-radius: 16px;
        padding: 24px 28px;
        margin-bottom: 24px;
    }
    .hero-title {
        color: #1E3A8A;
        font-size: 26px;
        font-weight: 700;
        margin: 0 0 6px 0;
    }
    .hero-subtitle {
        color: #2563EB;
        font-size: 14px;
        font-weight: 500;
        margin: 0;
    }

    /* Section Headings */
    .section-label {
        font-size: 14px;
        font-weight: 700;
        color: #1E293B;
        margin-bottom: 8px;
        display: flex;
        align-items: center;
        gap: 6px;
    }

    /* Row Badge */
    .row-badge {
        background-color: #2563EB;
        color: #FFFFFF;
        font-weight: 700;
        font-size: 13px;
        padding: 4px 14px;
        border-radius: 20px;
        display: inline-block;
        margin-bottom: 12px;
    }

    /* Status Badges */
    .status-ready {
        background-color: #DCFCE7;
        color: #15803D;
        font-weight: 600;
        font-size: 13px;
        padding: 8px 12px;
        border-radius: 8px;
        border: 1px solid #86EFAC;
        margin-bottom: 12px;
        text-align: center;
    }

    .status-wait {
        background-color: #FEF3C7;
        color: #B45309;
        font-weight: 600;
        font-size: 13px;
        padding: 8px 12px;
        border-radius: 8px;
        border: 1px solid #FDE68A;
        margin-bottom: 12px;
        text-align: center;
    }

    /* Order Item Box */
    .order-item-box {
        background-color: #F8FAFC;
        border: 1px solid #E2E8F0;
        border-radius: 8px;
        padding: 8px 14px;
        font-size: 13px;
        font-weight: 600;
        color: #1E293B;
        display: flex;
        align-items: center;
        margin-bottom: 6px;
    }

    /* Download Button Green Override */
    div[data-testid="stDownloadButton"] > button {
        background-color: #059669 !important;
        color: #FFFFFF !important;
        border: none !important;
        border-radius: 10px !important;
        font-weight: 600 !important;
        width: 100% !important;
        box-shadow: 0 2px 4px rgba(5, 150, 105, 0.2) !important;
    }
    
    div[data-testid="stDownloadButton"] > button:hover {
        background-color: #047857 !important;
        box-shadow: 0 4px 8px rgba(5, 150, 105, 0.3) !important;
    }

    /* Base Buttons Override */
    div.stButton > button {
        border-radius: 10px !important;
        font-weight: 600 !important;
        font-size: 14px !important;
    }
</style>
""", unsafe_allow_html=True)

# 3. Hero Header Component
st.markdown("""
<div class="hero-banner">
    <div class="hero-title">Anuan Ratu Kaylla dabel L</div>
    <div class="hero-subtitle">cobain bae dah</div>
</div>
""", unsafe_allow_html=True)

# Helper function to convert DOCX or PDF into PDF stream
def convert_file_to_pdf_stream(uploaded_file):
    fname = uploaded_file.name.lower()
    file_bytes = uploaded_file.read()
    uploaded_file.seek(0)
    
    if fname.endswith(".pdf"):
        return io.BytesIO(file_bytes)
    elif fname.endswith(".docx") or fname.endswith(".doc"):
        try:
            doc = Document(io.BytesIO(file_bytes))
            pdf_buffer = io.BytesIO()
            pdf_doc = SimpleDocTemplate(
                pdf_buffer,
                pagesize=letter,
                rightMargin=40, leftMargin=40,
                topMargin=40, bottomMargin=40
            )
            styles = getSampleStyleSheet()
            normal_style = styles['Normal']
            normal_style.fontSize = 11
            normal_style.leading = 14
            
            story = []
            for p in doc.paragraphs:
                text = p.text.strip()
                if text:
                    safe_text = text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                    story.append(Paragraph(safe_text, normal_style))
                    story.append(Spacer(1, 6))
                    
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        safe_row = row_text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                        story.append(Paragraph(f"<b>{safe_row}</b>", normal_style))
                        story.append(Spacer(1, 4))

            if not story:
                story.append(Paragraph("(Dokumen Word Kosong)", normal_style))
                
            pdf_doc.build(story)
            pdf_buffer.seek(0)
            return pdf_buffer
        except Exception as e:
            st.error(f"Gagal mengonversi file Word '{uploaded_file.name}': {e}")
            return None
    return None

# Helper function to merge ordered PDF streams
def merge_pdf_list(ordered_pdf_streams):
    merger = PyPDF2.PdfMerger()
    for pdf_stream in ordered_pdf_streams:
        if pdf_stream is not None:
            pdf_stream.seek(0)
            merger.append(pdf_stream)
            
    output = io.BytesIO()
    merger.write(output)
    merger.close()
    return output.getvalue()

# Helper function for HTML Base64 PDF Preview
def render_pdf_preview_iframe(pdf_bytes, height=450):
    base64_pdf = base64.b64encode(pdf_bytes).decode('utf-8')
    pdf_display = f'''
    <iframe 
        src="data:application/pdf;base64,{base64_pdf}#toolbar=1&navpanes=0&scrollbar=1" 
        width="100%" 
        height="{height}px" 
        type="application/pdf"
        style="border-radius: 12px; border: 1px solid #CBD5E1; background-color: #FFFFFF;"
    >
    </iframe>
    '''
    st.markdown(pdf_display, unsafe_allow_html=True)

# Session State Initialization
if 'num_rows' not in st.session_state:
    st.session_state.num_rows = 1

def add_row():
    st.session_state.num_rows += 1

def remove_row():
    if st.session_state.num_rows > 1:
        st.session_state.num_rows -= 1

def reset_rows():
    st.session_state.num_rows = 1

# List to collect processed data for master ZIP download
processed_files = []

# 4. Render Rows using Native Bordered Containers
for i in range(st.session_state.num_rows):
    with st.container(border=True):
        st.markdown(f'<span class="row-badge">Baris {i+1}</span>', unsafe_allow_html=True)
        
        c1, c2, c3 = st.columns([3.5, 4.5, 4])
        
        with c1:
            st.markdown('<div class="section-label">📄 File Utama (Bisa >1 PDF / Word)</div>', unsafe_allow_html=True)
            main_pdfs = st.file_uploader(
                "Upload File Utama", 
                type=["pdf", "docx", "doc"], 
                accept_multiple_files=True,
                key=f"main_{i}",
                label_visibility="collapsed"
            )
            
        with c2:
            st.markdown('<div class="section-label">📎 File Lampiran (Bisa >1 PDF / Word)</div>', unsafe_allow_html=True)
            attachments = st.file_uploader(
                "Upload Lampiran", 
                type=["pdf", "docx", "doc"], 
                accept_multiple_files=True, 
                key=f"att_{i}",
                label_visibility="collapsed"
            )
            
        with c3:
            st.markdown('<div class="section-label">💾 Nama File Output & Unduh</div>', unsafe_allow_html=True)
            default_name = f"hasil_gabungan_{i+1}.pdf"
            output_name = st.text_input(
                "Nama File Output",
                value=default_name,
                key=f"out_{i}",
                placeholder="contoh_file.pdf",
                label_visibility="collapsed"
            )
            
            clean_filename = output_name.strip()
            if not clean_filename.lower().endswith(".pdf"):
                clean_filename += ".pdf"

        # Gather all current files for Row i
        available_files = []
        if main_pdfs:
            for idx_main, m_file in enumerate(main_pdfs):
                icon = "📝 Word" if m_file.name.lower().endswith(('.docx', '.doc')) else "📄 PDF"
                available_files.append({"id": f"main_{idx_main}_{m_file.name}", "name": f"{icon} Utama {idx_main+1}: {m_file.name}", "obj": m_file})
            
        if attachments:
            for idx_att, att in enumerate(attachments):
                icon = "📝 Word" if att.name.lower().endswith(('.docx', '.doc')) else "📄 PDF"
                available_files.append({"id": f"att_{idx_att}_{att.name}", "name": f"{icon} Lampiran {idx_att+1}: {att.name}", "obj": att})

        # Reordering State Management per row
        order_key = f"file_order_{i}"
        
        if available_files:
            # Sync session state file order with uploaded files
            current_ids = [f["id"] for f in available_files]
            
            if order_key not in st.session_state:
                st.session_state[order_key] = current_ids
            else:
                # Remove deleted files & add newly uploaded files
                existing_order = [fid for fid in st.session_state[order_key] if fid in current_ids]
                for fid in current_ids:
                    if fid not in existing_order:
                        existing_order.append(fid)
                st.session_state[order_key] = existing_order
                
            # Sort files according to session state order
            id_to_file = {f["id"]: f for f in available_files}
            ordered_files_info = [id_to_file[fid] for fid in st.session_state[order_key] if fid in id_to_file]
            
            # --- REORDERING & PREVIEW SUB-SECTION ---
            st.markdown("---")
            reorder_col, preview_col = st.columns([1, 1])
            
            with reorder_col:
                st.markdown("##### 🔀 Atur Urutan File (PDF & Word)")
                st.caption("Gunakan tombol ⬆️ dan ⬇️ untuk mengubah posisi/urutan file sebelum digabungkan.")
                
                num_items = len(ordered_files_info)
                for idx, file_info in enumerate(ordered_files_info):
                    btn_c1, btn_c2, text_c = st.columns([0.8, 0.8, 5])
                    
                    with btn_c1:
                        if idx > 0:
                            if st.button("⬆️", key=f"up_{i}_{idx}", help="Naikkan posisi"):
                                st.session_state[order_key][idx], st.session_state[order_key][idx-1] = (
                                    st.session_state[order_key][idx-1], st.session_state[order_key][idx]
                                )
                                st.rerun()
                    
                    with btn_c2:
                        if idx < num_items - 1:
                            if st.button("⬇️", key=f"down_{i}_{idx}", help="Turunkan posisi"):
                                st.session_state[order_key][idx], st.session_state[order_key][idx+1] = (
                                    st.session_state[order_key][idx+1], st.session_state[order_key][idx]
                                )
                                st.rerun()
                                
                    with text_c:
                        st.markdown(f'<div class="order-item-box"><b>#{idx+1}</b>&nbsp; {file_info["name"]}</div>', unsafe_allow_html=True)
            
            # Convert files (both Word & PDF) to PDF streams and merge
            ordered_pdf_streams = []
            for f_info in ordered_files_info:
                pdf_stream = convert_file_to_pdf_stream(f_info["obj"])
                if pdf_stream:
                    ordered_pdf_streams.append(pdf_stream)
                    
            merged_data = merge_pdf_list(ordered_pdf_streams)
            
            # Page count summary
            total_pages = 0
            for pdf_stream in ordered_pdf_streams:
                pdf_stream.seek(0)
                r = PyPDF2.PdfReader(pdf_stream)
                total_pages += len(r.pages)
                
            with c3:
                st.markdown(f'<div class="status-ready">✓ Siap ({total_pages} Hal | {len(ordered_pdf_streams)} File)</div>', unsafe_allow_html=True)
                st.download_button(
                    label=f"📥 Unduh Hasil Baris {i+1}",
                    data=merged_data,
                    file_name=clean_filename,
                    mime="application/pdf",
                    key=f"dl_single_{i}"
                )
                
            with preview_col:
                st.markdown("##### 👁️ Preview Hasil Gabungan (Live)")
                with st.expander("Buka / Tutup Interactive PDF Viewer", expanded=True):
                    render_pdf_preview_iframe(merged_data, height=380)

            # Collect for ZIP download
            processed_files.append({
                "filename": clean_filename,
                "data": merged_data
            })
            
        else:
            with c3:
                st.markdown('<div class="status-wait">⏳ Upload File Utama untuk Memproses</div>', unsafe_allow_html=True)
                st.button(
                    f"📥 Unduh Hasil Baris {i+1}",
                    disabled=True,
                    key=f"dl_dis_{i}",
                    use_container_width=True
                )

# 5. Bottom Controls Container (Di Paling Bawah)
st.write("") # Spacer
with st.container(border=True):
    st.markdown("### ⚙️ Kontrol Baris & Unduh Masal")
    
    col_btn1, col_btn2, col_btn3 = st.columns([3, 2, 2])

    with col_btn1:
        st.button("➕ Tambah Baris Baru", on_click=add_row, use_container_width=True, type="primary")

    with col_btn2:
        st.button("➖ Hapus Baris Terakhir", on_click=remove_row, use_container_width=True, disabled=(st.session_state.num_rows <= 1))

    with col_btn3:
        st.button("🧹 Reset Semua Baris", on_click=reset_rows, use_container_width=True)

    st.divider()

    # Master ZIP Download Section
    if processed_files:
        zip_buffer = io.BytesIO()
        with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
            used_names = set()
            for item in processed_files:
                fname = item["filename"]
                # Resolve duplicate filename collision
                counter = 1
                original_fname = fname
                while fname in used_names:
                    name_part, ext_part = original_fname.rsplit('.', 1)
                    fname = f"{name_part}_({counter}).{ext_part}"
                    counter += 1
                used_names.add(fname)
                
                zip_file.writestr(fname, item["data"])

        st.success(f"🎉 **{len(processed_files)} dari {st.session_state.num_rows} Baris** berhasil digabungkan!")
        
        st.download_button(
            label=f"📦 Unduh Semua Sekaligus ({len(processed_files)} File .ZIP)",
            data=zip_buffer.getvalue(),
            file_name="semua_file_gabungan.zip",
            mime="application/zip",
            use_container_width=True
        )
    else:
        st.info("💡 Upload minimal 1 File Utama pada baris manapun untuk mengunduh seluruh file sekaligus dalam format .ZIP")
